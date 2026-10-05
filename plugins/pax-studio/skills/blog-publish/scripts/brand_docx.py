"""Brand styling helpers for article drafts written with python-docx (pip install python-docx).

Usage in an article builder (e.g. blog/make_ep01_docx.py):
    from docx import Document
    import brand_docx
    doc = Document(); brand_docx.apply(doc, brand_docx.load("config/blog.json"))
    ... doc.add_heading(...), doc.add_paragraph(...), t = doc.add_table(...); brand_docx.style_table(t, colors)
    doc.save("blog/final/ep01.docx")

Colors come from the "brand" block of the blog config: {"primary": "77372E", "accent": "7F417D",
"light": "E9E7E7", "text": "2B2421"} (hex without #).
"""
import json
from pathlib import Path
from docx.shared import RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DEFAULT = {"primary": "77372E", "accent": "7F417D", "light": "E9E7E7", "text": "2B2421"}


def load(config_path):
    cfg = json.loads(Path(config_path).read_text(encoding="utf8"))
    return {**DEFAULT, **{k: v.lstrip("#") for k, v in cfg.get("brand", {}).items()}}


def apply(doc, colors=DEFAULT):
    doc.styles["Normal"].font.color.rgb = RGBColor.from_string(colors["text"])
    for name, key in [("Title", "primary"), ("Heading 1", "primary"), ("Heading 2", "accent"), ("Heading 3", "accent")]:
        doc.styles[name].font.color.rgb = RGBColor.from_string(colors[key])


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def style_table(t, colors=DEFAULT):
    """Header row in the accent color with white bold text; zebra rows in the light color."""
    t.style = "Table Grid"
    for i, row in enumerate(t.rows):
        for c in row.cells:
            if i == 0:
                shade(c, colors["accent"])
                for r in c.paragraphs[0].runs:
                    r.font.color.rgb = RGBColor.from_string("FFFFFF"); r.bold = True
            elif i % 2 == 0:
                shade(c, colors["light"])
